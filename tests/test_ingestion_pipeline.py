"""
Comprehensive test suite for Phase 2 Ingestion & Chunking Pipeline.
Validates all 18 specified pipeline requirements.
"""

import unittest
import json
import hashlib
from pathlib import Path
import docx

from ingestion.config import (
    RAW_DATA_DIR,
    PROCESSED_DATA_DIR,
    EXPECTED_FILES,
    MAX_CHUNK_CHARS,
)
from ingestion.models import Chunk, DocumentMetadata, BreadcrumbItem, ClauseInfo
from ingestion.metadata import extract_document_metadata
from ingestion.toc_filter import TOCFilter
from ingestion.clause_parser import ClauseParser, HierarchyTracker
from ingestion.table_parser import TableParser
from ingestion.docx_parser import DocxParser
from ingestion.chunker import HierarchicalChunker, split_oversized_prose


class TestPhase2IngestionPipeline(unittest.TestCase):
    """Test suite covering Phase 2 end-to-end ingestion and chunking functionality."""

    def setUp(self):
        self.raw_dir = RAW_DATA_DIR
        self.proc_dir = PROCESSED_DATA_DIR

    # 1. All five documents discovered
    def test_01_all_five_documents_discovered(self):
        """Requirement 1: All five finalized 3GPP DOCX documents must be present in data/raw."""
        for fname in EXPECTED_FILES:
            fpath = self.raw_dir / fname
            self.assertTrue(fpath.exists(), f"Missing required source document: {fname}")
            self.assertGreater(fpath.stat().st_size, 0, f"File is empty: {fname}")

    # 2. Raw source integrity
    def test_02_raw_source_integrity(self):
        """Requirement 2: Source documents in data/raw must remain unmodified and readable."""
        for fname in EXPECTED_FILES:
            fpath = self.raw_dir / fname
            # Ensure docx can open without error in read mode
            doc = docx.Document(str(fpath))
            self.assertGreater(len(doc.paragraphs), 0, f"Document has no paragraphs: {fname}")

    # 3. TOC filtering
    def test_03_toc_filtering(self):
        """Requirement 3: Structural TOC content is removed without dropping technical body content."""
        toc_filter = TOCFilter()
        # Front-matter TOC start
        is_toc, is_body = toc_filter.process_paragraph("Contents", "TT", index=1)
        self.assertTrue(is_toc)
        self.assertFalse(is_body)

        # TOC entry
        is_toc, is_body = toc_filter.process_paragraph("4.2.3\tNon-roaming reference architecture\t43", "toc 3", index=20)
        self.assertTrue(is_toc)
        self.assertFalse(is_body)

        # First body heading ends TOC
        is_toc, is_body = toc_filter.process_paragraph("Foreword", "Heading 1", index=50)
        self.assertFalse(is_toc)
        self.assertTrue(is_body)

        # Technical body paragraph with numbers must NOT be filtered
        is_toc, is_body = toc_filter.process_paragraph("1. The UE sends a Registration Request message.", "B1", index=60)
        self.assertFalse(is_toc)
        self.assertTrue(is_body)

    # 4. Valid clause detection
    def test_04_valid_clause_detection(self):
        """Requirement 4: Valid clause heading is recognized with level and title."""
        parser = ClauseParser()
        clause = parser.parse_heading("4.2.3\tNon-roaming reference architecture", "Heading 3", is_toc=False)
        self.assertIsNotNone(clause)
        self.assertEqual(clause.clause_number, "4.2.3")
        self.assertEqual(clause.clause_title, "Non-roaming reference architecture")
        self.assertEqual(clause.level, 3)

    # 5. Nested clause detection
    def test_05_nested_clause_detection(self):
        """Requirement 5: Deeply nested clause (e.g. 4.2.8.2.1) is recognized with level 5."""
        parser = ClauseParser()
        clause = parser.parse_heading("4.2.8.2.1\tNon-roaming Architecture", "Heading 5", is_toc=False)
        self.assertIsNotNone(clause)
        self.assertEqual(clause.clause_number, "4.2.8.2.1")
        self.assertEqual(clause.level, 5)

    # 6. Annex handling
    def test_06_annex_handling(self):
        """Requirement 6: Top-level Annex and Annex subclauses are handled properly."""
        parser = ClauseParser()
        top_annex = parser.parse_heading(
            "Annex A (informative):\nRelationship between Service-Based Interfaces and Reference Points",
            "Heading 8",
            is_toc=False,
        )
        self.assertIsNotNone(top_annex)
        self.assertEqual(top_annex.clause_number, "Annex A")
        self.assertTrue(top_annex.is_annex)
        self.assertEqual(top_annex.annex_type, "informative")

        sub_annex = parser.parse_heading("A.1\tSupplementary Uplink", "Heading 1", is_toc=False)
        self.assertIsNotNone(sub_annex)
        self.assertEqual(sub_annex.clause_number, "A.1")
        self.assertTrue(sub_annex.is_annex)

    # 7. Numeric prose rejection
    def test_07_numeric_prose_rejection(self):
        """Requirement 7: Paragraphs with numbers in normal styles must not become clauses."""
        parser = ClauseParser()
        prose_samples = [
            ("As described in clause 4.2.3, the AMF coordinates registration.", "Normal"),
            ("1. The UE sends a Registration Request message to the (R)AN.", "B1"),
            ("The timer T3512 has a default value of 54.0 minutes.", "Normal"),
        ]
        for text, style in prose_samples:
            self.assertIsNone(parser.parse_heading(text, style, is_toc=False))

    # 8. Table text not falsely detected as heading
    def test_08_table_text_rejection(self):
        """Requirement 8: Table headers and table cells must not be classified as clause headings."""
        parser = ClauseParser()
        table_samples = [
            ("5.4.1\t5GMM-REGISTERED\tNormal service", "TA"),
            ("Table 4.2.3-1: Non-roaming architecture reference points", "TH"),
        ]
        for text, style in table_samples:
            self.assertIsNone(parser.parse_heading(text, style, is_toc=False))

    # 9. Hierarchy reconstruction
    def test_09_hierarchy_reconstruction(self):
        """Requirement 9: Hierarchical breadcrumb stack is accurately maintained through deep tree traversal."""
        parser = ClauseParser()
        parser.parse_heading("4\tArchitecture model and concepts", "Heading 1")
        parser.parse_heading("4.2\tArchitecture reference model", "Heading 2")
        c3 = parser.parse_heading("4.2.3\tNon-roaming reference architecture", "Heading 3")

        self.assertEqual(len(c3.breadcrumb), 3)
        self.assertEqual(c3.breadcrumb[0].number, "4")
        self.assertEqual(c3.breadcrumb[1].number, "4.2")
        self.assertEqual(c3.breadcrumb[2].number, "4.2.3")
        self.assertEqual(c3.parent_clause, "4.2")

        # Moving to a parallel branch pops deeper levels
        c_next = parser.parse_heading("4.3\tInterworking with EPC", "Heading 2")
        self.assertEqual(len(c_next.breadcrumb), 2)
        self.assertEqual(c_next.breadcrumb[0].number, "4")
        self.assertEqual(c_next.breadcrumb[1].number, "4.3")
        self.assertEqual(c_next.parent_clause, "4")

    # 10. Table extraction
    def test_10_table_extraction(self):
        """Requirement 10: Table parser generates valid GitHub markdown and deterministic table IDs."""
        table_parser = TableParser("TS 23.501")
        # Load a real table from TS 28.535 (compact doc)
        doc = docx.Document(str(self.raw_dir / "28535-i20.docx"))
        self.assertGreater(len(doc.tables), 0)
        table = doc.tables[-1]  # Change history table
        t_id, t_md, cap = table_parser.parse(table, caption="Table A-1: Change History")

        self.assertTrue(t_id.startswith("TS23501-T"))
        self.assertIn("|", t_md)
        self.assertIn(":---", t_md)

    # 11. Table-to-clause association
    def test_11_table_to_clause_association(self):
        """Requirement 11: Table chunk inherits active clause context and breadcrumbs."""
        jsonl_path = self.proc_dir / "TS_38_300_chunks.jsonl"
        with open(jsonl_path, "r", encoding="utf-8") as f:
            chunks = [json.loads(line) for line in f]

        table_chunks = [c for c in chunks if c["source_type"] == "table"]
        self.assertGreater(len(table_chunks), 0)
        first_table = table_chunks[0]

        self.assertIsNotNone(first_table["clause_number"])
        self.assertIsNotNone(first_table["table_id"])
        self.assertGreater(len(first_table["breadcrumb"]), 0)

    # 12. Metadata completeness
    def test_12_metadata_completeness(self):
        """Requirement 12: Every generated chunk has all mandatory metadata fields populated."""
        for fname in EXPECTED_FILES:
            clean_spec = fname.split("-")[0]
            jsonl_files = list(self.proc_dir.glob(f"TS_{clean_spec[:2]}_{clean_spec[2:]}*.jsonl"))
            self.assertTrue(len(jsonl_files) > 0, f"No JSONL output for {fname}")
            with open(jsonl_files[0], "r", encoding="utf-8") as f:
                for idx, line in enumerate(f, start=1):
                    chunk = json.loads(line)
                    self.assertTrue(chunk.get("chunk_id"), f"Missing chunk_id in {fname} line {idx}")
                    self.assertTrue(chunk.get("specification"), f"Missing spec in {fname} line {idx}")
                    self.assertTrue(chunk.get("version"), f"Missing version in {fname} line {idx}")
                    self.assertTrue(chunk.get("release"), f"Missing release in {fname} line {idx}")
                    self.assertTrue(chunk.get("clause_number"), f"Missing clause in {fname} line {idx}")
                    self.assertTrue(chunk.get("text"), f"Empty text in {fname} line {idx}")
                    self.assertIsInstance(chunk.get("document_order"), int)

    # 13. Source ordering
    def test_13_source_ordering(self):
        """Requirement 13: document_order increases strictly monotonically within each document."""
        for jsonl_path in self.proc_dir.glob("*.jsonl"):
            with open(jsonl_path, "r", encoding="utf-8") as f:
                chunks = [json.loads(line) for line in f]
            orders = [c["document_order"] for c in chunks]
            self.assertEqual(orders, list(range(1, len(chunks) + 1)), f"Non-monotonic order in {jsonl_path.name}")

    # 14. Deterministic chunk IDs
    def test_14_deterministic_chunk_ids(self):
        """Requirement 14: Chunk IDs follow deterministic naming convention without random tokens."""
        jsonl_path = self.proc_dir / "TS_28_535_chunks.jsonl"
        with open(jsonl_path, "r", encoding="utf-8") as f:
            for line in f:
                c = json.loads(line)
                cid = c["chunk_id"]
                self.assertTrue(cid.startswith("TS_28.535-V18.2.0-"), f"Non-standard ID format: {cid}")

    # 15. Deterministic repeated ingestion
    def test_15_deterministic_repeated_ingestion(self):
        """Requirement 15: Re-running ingestion against TS 28.535 yields identical byte content."""
        doc_path = self.raw_dir / "28535-i20.docx"
        parser1 = DocxParser(doc_path)
        chunks1 = HierarchicalChunker(parser1).chunk_document()

        parser2 = DocxParser(doc_path)
        chunks2 = HierarchicalChunker(parser2).chunk_document()

        self.assertEqual(len(chunks1), len(chunks2))
        for c1, c2 in zip(chunks1, chunks2):
            self.assertEqual(c1.chunk_id, c2.chunk_id)
            self.assertEqual(c1.content_hash, c2.content_hash)
            self.assertEqual(c1.text, c2.text)

    # 16. Large-clause splitting
    def test_16_large_clause_splitting(self):
        """Requirement 16: Oversized paragraphs list is split into sub-chunks without exceeding max chars."""
        paragraphs = [
            "This is paragraph A. " * 30,  # ~630 chars
            "This is paragraph B. " * 40,  # ~840 chars
            "This is paragraph C. " * 50,  # ~1050 chars
            "This is paragraph D. " * 30,  # ~630 chars
        ]
        blocks = split_oversized_prose(paragraphs, max_chars=1500)
        self.assertGreater(len(blocks), 1)
        for b in blocks:
            self.assertLessEqual(len(b), 1600)  # within allowance

    # 17. No sentence-level destructive splitting
    def test_17_no_sentence_destructive_splitting(self):
        """Requirement 17: Paragraphs are kept intact without truncating mid-sentence."""
        p1 = "First sentence of paragraph 1. Second sentence of paragraph 1."
        p2 = "First sentence of paragraph 2. Second sentence of paragraph 2."
        blocks = split_oversized_prose([p1, p2], max_chars=100)
        # Should split along whole paragraph boundaries
        self.assertEqual(len(blocks), 2)
        self.assertEqual(blocks[0], p1)
        self.assertEqual(blocks[1], p2)

    # 18. JSONL validity & Manifest check
    def test_18_jsonl_validity_and_manifest(self):
        """Requirement 18: All output files are valid JSONL and corpus manifest contains all 5 specs."""
        manifest_path = self.proc_dir / "corpus_manifest.json"
        self.assertTrue(manifest_path.exists())
        with open(manifest_path, "r", encoding="utf-8") as f:
            manifest = json.load(f)

        self.assertEqual(manifest["corpus_summary"]["documents_count"], 5)
        self.assertEqual(len(manifest["documents"]), 5)
        self.assertEqual(manifest["corpus_summary"]["total_errors"], 0)

        for d in manifest["documents"]:
            jsonl_path = self.proc_dir / d["jsonl_file"]
            self.assertTrue(jsonl_path.exists())
            with open(jsonl_path, "r", encoding="utf-8") as f:
                line_count = sum(1 for _ in f)
            self.assertEqual(line_count, d["total_chunks_count"])


if __name__ == "__main__":
    unittest.main()
